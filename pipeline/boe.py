"""Official Bank of England nominal/real gilt spot curves and RPI breakevens.

Recognises the two whole-curve sheet names inspected in the official archive.
Missing tenor observations remain null; entirely empty non-trading dates are
excluded. No interpolation is performed.
"""
import io
import math
import re
import zipfile
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

PAGE_URL = 'https://www.bankofengland.co.uk/statistics/yield-curves'
LATEST_URL = 'https://www.bankofengland.co.uk/-/media/boe/files/statistics/yield-curves/latest-yield-curve-data.zip'
REAL_ARCHIVE_URL = 'https://www.bankofengland.co.uk/-/media/boe/files/statistics/yield-curves/glcrealddata.zip'
ARCHIVE_URL = 'https://www.bankofengland.co.uk/-/media/boe/files/statistics/yield-curves/glcnominalddata.zip'


def parse_zip(raw, kind="nominal"):
    """Read the selected whole gilt spot curve; never forward or OIS."""
    if kind not in ("nominal", "real"):
        raise ValueError("Unsupported gilt curve kind")
    curves = {}
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        names = [n for n in archive.namelist() if f'glc {kind}' in n.lower() and n.lower().endswith('.xlsx')]
        if not names:
            raise ValueError(f'No recognised {kind} gilt workbook in BoE ZIP')
        for name in names:
            if archive.getinfo(name).file_size > 100_000_000:
                raise ValueError('BoE workbook exceeds import size limit')
            workbook = load_workbook(io.BytesIO(archive.read(name)), data_only=True, read_only=True)
            try:
                sheets = [s for s in workbook.sheetnames if re.fullmatch(rf'4\.\s*(?:{kind} )?spot curve', s.strip(), re.I)]
                if len(sheets) != 1:
                    raise ValueError('Expected exactly one whole spot-curve worksheet')
                rows = list(workbook[sheets[0]].values)
                headers = []
                for index, row in enumerate(rows[:10]):
                    if str(row[0]).strip().lower().rstrip(':') != 'years':
                        continue
                    numbers = [(i, float(v)) for i, v in enumerate(row[1:], 1) if isinstance(v, (int, float)) and not isinstance(v, bool)]
                    if len(numbers) >= 3 and all(0 < v <= 40 for _, v in numbers):
                        headers.append((index, numbers))
                if len(headers) != 1:
                    raise ValueError('Expected unique numeric maturity header in years')
                header_index, tenors = headers[0]
                if any(b[1] <= a[1] for a, b in zip(tenors, tenors[1:])):
                    raise ValueError('Tenors must be strictly increasing')
                for row in rows[header_index + 1:]:
                    if not isinstance(row[0], (date, datetime)):
                        continue
                    observed = row[0].date() if isinstance(row[0], datetime) else row[0]
                    points = []
                    for column, tenor in tenors:
                        value = row[column] if column < len(row) else None
                        if value is None or value == '':
                            points.append({'tenor': tenor, 'rate': None})
                            continue
                        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or abs(value) > 100:
                            raise ValueError('Invalid spot rate in BoE workbook')
                        points.append({'tenor': tenor, 'rate': float(value)})
                    if any(p['rate'] is not None for p in points):
                        key = observed.isoformat()
                        if key in curves and curves[key] != points:
                            raise ValueError('Conflicting dates within BoE archive')
                        curves[key] = points
            finally:
                workbook.close()
    if not curves:
        raise ValueError(f'No dated {kind} spot rates in BoE ZIP')
    return curves


def derive_breakevens(nominal, real):
    """Subtract matched continuously compounded spots; keep missing tenors null.

    No nearest-date substitution or tenor interpolation is permissible. Output
    uses the union of each matched date's tenor grids to expose coverage gaps.
    """
    result = {}
    for observed in sorted(nominal.keys() & real.keys()):
        n = {p['tenor']: p['rate'] for p in nominal[observed]}
        r = {p['tenor']: p['rate'] for p in real[observed]}
        points = [
            {'tenor': t, 'rate': n[t] - r[t] if n.get(t) is not None and r.get(t) is not None else None}
            for t in sorted(n.keys() | r.keys())
        ]
        if any(p['rate'] is not None for p in points):
            result[observed] = points
    if not result:
        raise ValueError('No matched nominal and real spot observations')
    return result


def fetch_curve(fetch):
    # The latest-month file takes precedence over an older archived vintage.
    nominal = parse_zip(fetch(ARCHIVE_URL))
    real = parse_zip(fetch(REAL_ARCHIVE_URL), 'real')
    latest_raw = fetch(LATEST_URL)
    nominal.update(parse_zip(latest_raw))
    real.update(parse_zip(latest_raw, 'real'))
    breakevens = derive_breakevens(nominal, real)
    latest = date.fromisoformat(max(nominal))
    cutoff = (latest - timedelta(days=370)).isoformat()
    def rows(curves):
        return [{'date': d, 'points': curves[d]} for d in sorted(curves) if d >= cutoff]
    if not rows(breakevens):
        raise ValueError('No matched breakevens in retained date range')
    return {
        'schemaVersion': 1, 'status': 'available', 'units': '% per annum',
        'basis': 'Bank of England fitted gilt zero-coupon spot curves, continuously compounded; tenor in years. RPI breakeven = nominal minus real on the same observation date and maturity.',
        'asOf': latest.isoformat(), 'observationDate': latest.isoformat(), 'publicationDate': None,
        'realAsOf': max(real), 'breakevenAsOf': max(breakevens),
        'sources': [{'id': 'boe-curve', 'publisher': 'Bank of England', 'title': 'Nominal and real gilt spot yield curves; derived RPI breakevens', 'url': PAGE_URL, 'downloadUrls': [ARCHIVE_URL, REAL_ARCHIVE_URL, LATEST_URL]}],
        'curves': rows(nominal), 'realCurves': rows(real), 'breakevenCurves': rows(breakevens),
        'notes': [
            'Missing tenor values remain null without interpolation. Entirely empty non-trading dates are excluded. Observation date is not the publication or retrieval date.',
            'Breakevens are market inflation compensation, including inflation risk premia and relative liquidity effects, not pure inflation expectations or CPI forecasts.',
            'Index-linked gilts reference RPI. The February 2030 alignment of RPI methods and data sources with CPIH can affect interpretation across that horizon; the Bank has not adjusted its fitting methodology for this change.',
        ],
    }
