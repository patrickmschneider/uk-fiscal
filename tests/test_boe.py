"""Parser safeguards plus optional local official-workbook regression.

The official fixture is local-only pending redistribution permission.
"""
import io
import unittest
import zipfile
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from pipeline.boe import parse_zip, derive_breakevens, fetch_curve, ARCHIVE_URL, REAL_ARCHIVE_URL, LATEST_URL


def synthetic_zip(sheet='4. spot curve', filename='GLC Nominal daily data.xlsx', rates=(4.1, None, -0.2)):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(['Synthetic parser test fixture — not real observations'])
    ws.append(['years', 0.5, 1, 2])
    ws.append([datetime(2026, 1, 5), *rates])
    excel = io.BytesIO()
    wb.save(excel)
    result = io.BytesIO()
    with zipfile.ZipFile(result, 'w') as z:
        z.writestr(filename, excel.getvalue())
    return result.getvalue()


class CurveTests(unittest.TestCase):
    def test_missing_tenors_not_filled_and_negative_rates_valid(self):
        result = parse_zip(synthetic_zip())
        self.assertEqual(result['2026-01-05'], [{'tenor': 0.5, 'rate': 4.1}, {'tenor': 1.0, 'rate': None}, {'tenor': 2.0, 'rate': -0.2}])

    def test_legacy_nominal_sheet_name(self):
        self.assertEqual(parse_zip(synthetic_zip()), parse_zip(synthetic_zip(sheet='4. nominal spot curve')))

    def test_empty_non_trading_dates_excluded(self):
        with self.assertRaisesRegex(ValueError, 'No dated nominal spot rates'):
            parse_zip(synthetic_zip(rates=(None, None, None)))

    def test_official_current_workbook(self):
        fixture = Path(__file__).resolve().parents[1] / 'data/fixtures/boe/2026-09-06-latest-yield-curve-data.zip'
        if not fixture.exists():
            self.skipTest('Local official fixture absent; download privately to validate real source')
        curves = parse_zip(fixture.read_bytes())
        self.assertEqual(sorted(curves), ['2026-09-01', '2026-09-02', '2026-09-03'])
        self.assertEqual(len(curves['2026-09-03']), 80)
        self.assertEqual(curves['2026-09-03'][19]['tenor'], 10.0)
        self.assertAlmostEqual(curves['2026-09-03'][19]['rate'], 5.1575084015148365)
        self.assertAlmostEqual(curves['2026-09-03'][-1]['rate'], 5.619299573940564)

    def test_rejects_other_curves(self):
        for raw in [synthetic_zip(sheet='5. forward curve'), synthetic_zip(filename='GLC Real daily data.xlsx')]:
            with self.assertRaises(ValueError):
                parse_zip(raw)

    def test_rejects_non_numeric_rates(self):
        with self.assertRaises(ValueError):
            parse_zip(synthetic_zip(rates=(4.1, 'changed layout', 2)))

    def test_real_spot_selection(self):
        real = synthetic_zip(filename='GLC Real daily data.xlsx', sheet='4. real spot curve')
        self.assertEqual(parse_zip(real, 'real')['2026-01-05'][0]['rate'], 4.1)
        with self.assertRaises(ValueError):
            parse_zip(synthetic_zip(), 'real')

    def test_breakevens_require_exact_date_and_tenor_match(self):
        nominal = {
            '2026-01-05': [{'tenor': 1, 'rate': 4}, {'tenor': 2, 'rate': 5}, {'tenor': 3, 'rate': None}],
            '2026-01-06': [{'tenor': 1, 'rate': 4}],
        }
        real = {
            '2026-01-05': [{'tenor': 2, 'rate': -0.5}, {'tenor': 3, 'rate': 1}, {'tenor': 4, 'rate': 2}],
            '2026-01-07': [{'tenor': 1, 'rate': 1}],
        }
        self.assertEqual(derive_breakevens(nominal, real), {
            '2026-01-05': [{'tenor': 1, 'rate': None}, {'tenor': 2, 'rate': 5.5}, {'tenor': 3, 'rate': None}, {'tenor': 4, 'rate': None}],
        })

    def test_no_overlap_fails_instead_of_fabricating_curve(self):
        for real in ({'2026-01-06': [{'tenor': 1, 'rate': 1}]}, {'2026-01-05': [{'tenor': 2, 'rate': 1}]}):
            with self.assertRaisesRegex(ValueError, 'No matched'):
                derive_breakevens({'2026-01-05': [{'tenor': 1, 'rate': 4}]}, real)

    def test_official_breakeven_matches_published_inflation_workbook(self):
        fixture = Path(__file__).resolve().parents[1] / 'data/fixtures/boe/2026-09-06-latest-yield-curve-data.zip'
        if not fixture.exists():
            self.skipTest('Local official fixture absent; download privately to validate real source')
        from openpyxl import load_workbook
        raw = fixture.read_bytes()
        calculated = derive_breakevens(parse_zip(raw), parse_zip(raw, 'real'))
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            name = next(n for n in z.namelist() if 'GLC Inflation' in n)
            workbook = load_workbook(io.BytesIO(z.read(name)), read_only=True, data_only=True)
            try:
                rows = list(workbook['4. spot curve'].values)
            finally:
                workbook.close()
        tenors = next(row for row in rows if str(row[0]).lower().rstrip(':') == 'years')
        checked = 0
        for row in rows:
            if not isinstance(row[0], datetime):
                continue
            result = {p['tenor']: p['rate'] for p in calculated.get(row[0].date().isoformat(), [])}
            for t, rate in zip(tenors[1:], row[1:]):
                if isinstance(rate, (int, float)):
                    self.assertAlmostEqual(result[t], rate, places=10)
                    checked += 1
        self.assertGreater(checked, 100)

    def test_latest_month_overrides_each_archive_before_matching(self):
        combined = io.BytesIO()
        with zipfile.ZipFile(combined, 'w') as latest:
            for raw in [synthetic_zip(rates=(5, 6, 7)), synthetic_zip(filename='GLC Real daily data.xlsx', rates=(1, None, 2))]:
                with zipfile.ZipFile(io.BytesIO(raw)) as part:
                    for name in part.namelist():
                        latest.writestr(name, part.read(name))
        inputs = {
            ARCHIVE_URL: synthetic_zip(rates=(10, 11, 12)),
            REAL_ARCHIVE_URL: synthetic_zip(filename='GLC Real daily data.xlsx', rates=(4, 5, 6)),
            LATEST_URL: combined.getvalue(),
        }
        result = fetch_curve(inputs.__getitem__)
        self.assertEqual(result['curves'][0]['points'][0]['rate'], 5)
        self.assertEqual(result['realCurves'][0]['points'][0]['rate'], 1)
        self.assertEqual([p['rate'] for p in result['breakevenCurves'][0]['points']], [4, None, 5])
